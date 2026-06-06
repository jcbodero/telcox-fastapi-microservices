from __future__ import annotations

import html
import os
import re
import sys
import textwrap
import xml.etree.ElementTree as ET
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
DEPS = ROOT / ".generated" / "docx_deps"
if DEPS.exists():
    sys.path.insert(0, str(DEPS))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


MD_PATH = ROOT / "docs" / "ARCHITECTURE_ONPREM_AWS.md"
DRAWIO_PATH = ROOT / "docs" / "TELCOX_ARCHITECTURE.drawio"
OUT_DIR = ROOT / "docs" / "generated_report_assets"
OUT_DOCX = ROOT / "docs" / "TelcoX_Informe_Arquitectura_Profesional.docx"


NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(90, 96, 106)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf") if bold else Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf") if bold else Path("C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def clean_label(value: str | None) -> str:
    if not value:
        return ""
    value = html.unescape(value)
    value = re.sub(r"<[^>]+>", "", value)
    value = value.replace("\r", "")
    return value


def wrap_text(draw: ImageDraw.ImageDraw, text: str, max_width: int, fnt: ImageFont.ImageFont) -> list[str]:
    lines: list[str] = []
    for raw in text.split("\n"):
        words = raw.split()
        if not words:
            lines.append("")
            continue
        current = words[0]
        for word in words[1:]:
            test = current + " " + word
            if draw.textbbox((0, 0), test, font=fnt)[2] <= max_width:
                current = test
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return lines


def style_colors(style: str) -> tuple[str, str]:
    fill = "#FFFFFF"
    stroke = "#4B5563"
    for token in style.split(";"):
        if token.startswith("fillColor="):
            fill = token.split("=", 1)[1] or fill
        elif token.startswith("strokeColor="):
            stroke = token.split("=", 1)[1] or stroke
    def normalize(value: str) -> str:
        if value.startswith("light-dark("):
            inner = value[len("light-dark(") :].rstrip(")")
            return inner.split(",", 1)[0].strip()
        if not value.startswith("#") and re.fullmatch(r"[0-9A-Fa-f]{6}", value):
            return f"#{value}"
        return value

    fill = normalize(fill)
    stroke = normalize(stroke)
    if fill == "none":
        fill = "#FFFFFF"
    if stroke == "none":
        stroke = "#FFFFFF"
    return fill, stroke


def render_drawio_pages() -> list[tuple[str, Path]]:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    tree = ET.parse(DRAWIO_PATH)
    root = tree.getroot()
    outputs: list[tuple[str, Path]] = []

    for diagram in root.findall("diagram"):
        page_name = diagram.attrib.get("name", "Diagrama")
        if page_name == "On-Premise":
            outputs.append((
                page_name,
                render_flow_diagram(
                    "On-Premise",
                    [
                        ("U", "Usuarios web y movil", 45, 285),
                        ("W", "Cloudflare WAF\nDominio publico\nSSL/TLS", 265, 285),
                        ("G", "Ingress / API Gateway\nTraefik", 505, 285),
                        ("K", "Keycloak\nOIDC + PKCE", 505, 425),
                        ("UI", "Frontend Web\nNext.js", 745, 120),
                        ("M", "App Movil\nExpo / React Native", 745, 245),
                        ("S", "Kubernetes Pods\nMicroservicios FastAPI", 745, 390),
                        ("A", "audit_service\nTrazabilidad", 1015, 120),
                        ("P", "PostgreSQL\nCifrado en reposo", 1015, 255),
                        ("N", "MongoDB", 1015, 390),
                        ("L", "Loki / Grafana\nObservabilidad", 1015, 525),
                    ],
                    [("U", "W"), ("W", "G"), ("G", "K"), ("G", "UI"), ("G", "M"), ("G", "S"), ("S", "A"), ("S", "P"), ("S", "N"), ("S", "L")],
                    OUT_DIR / "drawio_on_premise.png",
                    prefix="Draw.io",
                ),
            ))
            continue
        if page_name == "AWS Objetivo":
            outputs.append((
                page_name,
                render_flow_diagram(
                    "AWS Objetivo",
                    [
                        ("U", "Usuarios web y movil", 60, 260),
                        ("E", "CloudFront / ALB\nAWS WAF + ACM TLS", 300, 260),
                        ("ID", "Keycloak / OIDC", 555, 110),
                        ("C", "ECS Fargate / EKS\nMicroservicios FastAPI", 555, 260),
                        ("R", "RDS PostgreSQL\nMulti-AZ", 820, 110),
                        ("X", "ElastiCache Redis", 820, 260),
                        ("B", "Amazon S3", 820, 410),
                        ("W", "CloudWatch / Grafana", 1085, 260),
                        ("S", "IAM + Secrets Manager", 555, 535),
                    ],
                    [("U", "E"), ("E", "ID"), ("E", "C"), ("C", "R"), ("C", "X"), ("C", "B"), ("C", "W"), ("S", "C")],
                    OUT_DIR / "drawio_aws_objetivo.png",
                    prefix="Draw.io",
                ),
            ))
            continue
        model = diagram.find("mxGraphModel")
        if model is None:
            model = ET.fromstring(diagram.text or "")
        cells = model.find("root")
        if cells is None:
            continue

        vertices: dict[str, dict] = {}
        edges: list[dict] = []
        parent_offsets = {"1": (0.0, 0.0)}

        for cell in cells.findall("mxCell"):
            cell_id = cell.attrib.get("id")
            if not cell_id:
                continue
            geom = cell.find("mxGeometry")
            if cell.attrib.get("vertex") == "1" and geom is not None:
                parent = cell.attrib.get("parent", "1")
                px, py = parent_offsets.get(parent, (0.0, 0.0))
                x = float(geom.attrib.get("x", "0")) + px
                y = float(geom.attrib.get("y", "0")) + py
                w = float(geom.attrib.get("width", "120"))
                h = float(geom.attrib.get("height", "60"))
                label = clean_label(cell.attrib.get("value"))
                if not label:
                    continue
                vertices[cell_id] = {
                    "id": cell_id,
                    "label": label,
                    "style": cell.attrib.get("style", ""),
                    "x": x,
                    "y": y,
                    "w": w,
                    "h": h,
                }
                parent_offsets[cell_id] = (x, y)
            elif cell.attrib.get("edge") == "1":
                edges.append(cell.attrib)

        if not vertices:
            continue

        min_x = min(v["x"] for v in vertices.values()) - 35
        min_y = min(v["y"] for v in vertices.values()) - 35
        max_x = max(v["x"] + v["w"] for v in vertices.values()) + 35
        max_y = max(v["y"] + v["h"] for v in vertices.values()) + 35
        scale = min(1500 / (max_x - min_x), 950 / (max_y - min_y))
        scale = max(0.65, min(scale, 1.15))
        top_band = 58
        width = int((max_x - min_x) * scale)
        height = int((max_y - min_y) * scale) + top_band

        img = Image.new("RGB", (width, height), "#FFFFFF")
        draw = ImageDraw.Draw(img)

        def pt(v: dict, ax: float = 0.5, ay: float = 0.5) -> tuple[int, int]:
            return (
                int((v["x"] + v["w"] * ax - min_x) * scale),
                int((v["y"] + v["h"] * ay - min_y) * scale) + top_band,
            )

        containers = [v for v in vertices.values() if "swimlane" in v["style"]]
        normals = [v for v in vertices.values() if v not in containers and v["id"] not in {"op-title", "aws-title", "ci-title", "db-title"}]

        for v in containers:
            x1 = int((v["x"] - min_x) * scale)
            y1 = int((v["y"] - min_y) * scale) + top_band
            x2 = int((v["x"] + v["w"] - min_x) * scale)
            y2 = int((v["y"] + v["h"] - min_y) * scale) + top_band
            fill, stroke = style_colors(v["style"])
            draw.rounded_rectangle((x1, y1, x2, y2), radius=10, fill=fill, outline=stroke, width=2)
            draw.rectangle((x1, y1, x2, y1 + 34), fill="#F8FAFC", outline=stroke, width=1)
            draw.text((x1 + 12, y1 + 8), v["label"], fill="#111827", font=font(15, True))

        for edge in edges:
            source = vertices.get(edge.get("source", ""))
            target = vertices.get(edge.get("target", ""))
            if not source or not target:
                continue
            sx, sy = pt(source, 1.0, 0.5)
            tx, ty = pt(target, 0.0, 0.5)
            color = "#4B5563"
            draw.line((sx, sy, tx, ty), fill=color, width=2)
            if tx >= sx:
                arrow = [(tx, ty), (tx - 9, ty - 5), (tx - 9, ty + 5)]
            else:
                arrow = [(tx, ty), (tx + 9, ty - 5), (tx + 9, ty + 5)]
            draw.polygon(arrow, fill=color)

        for v in normals:
            x1 = int((v["x"] - min_x) * scale)
            y1 = int((v["y"] - min_y) * scale) + top_band
            x2 = int((v["x"] + v["w"] - min_x) * scale)
            y2 = int((v["y"] + v["h"] - min_y) * scale) + top_band
            fill, stroke = style_colors(v["style"])
            draw.rounded_rectangle((x1, y1, x2, y2), radius=10, fill=fill, outline=stroke, width=2)
            label = v["label"]
            fnt = font(15 if len(label) < 90 else 13, bold=True)
            lines = wrap_text(draw, label, max(40, x2 - x1 - 18), fnt)
            line_h = int(fnt.size * 1.22) if hasattr(fnt, "size") else 16
            total_h = min(len(lines) * line_h, y2 - y1 - 12)
            ty = y1 + max(8, (y2 - y1 - total_h) // 2)
            text_fill = "#FFFFFF" if fill.lower() in {"#0b2545", "#0969da", "#1d63ed", "#0f1689", "#24292f", "#326ce5"} else "#111827"
            for line in lines[: max(1, (y2 - y1 - 12) // line_h)]:
                bbox = draw.textbbox((0, 0), line, font=fnt)
                tx = x1 + max(8, (x2 - x1 - (bbox[2] - bbox[0])) // 2)
                draw.text((tx, ty), line, fill=text_fill, font=fnt)
                ty += line_h

        title = f"Draw.io - {page_name}"
        draw.rectangle((0, 0, width, 42), fill="#0B2545")
        draw.text((22, 10), title, fill="#FFFFFF", font=font(20, bold=True))

        out = OUT_DIR / f"drawio_{page_name.lower().replace(' ', '_').replace('-', '_')}.png"
        img.save(out, "PNG")
        outputs.append((page_name, out))
    return outputs


def render_flow_diagram(name: str, nodes: list[tuple[str, str, int, int]], edges: list[tuple[str, str]], out: Path, prefix: str = "Mermaid") -> Path:
    img = Image.new("RGB", (1400, 720), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, 1400, 52), fill="#0B2545")
    draw.text((26, 13), f"{prefix} - {name}", fill="#FFFFFF", font=font(22, bold=True))
    by_id = {}
    for node_id, label, x, y in nodes:
        by_id[node_id] = (x, y, 190, 82)
    for a, b in edges:
        ax, ay, aw, ah = by_id[a]
        bx, by, bw, bh = by_id[b]
        sx, sy = ax + aw, ay + ah // 2
        tx, ty = bx, by + bh // 2
        draw.line((sx, sy, tx, ty), fill="#4B5563", width=3)
        draw.polygon([(tx, ty), (tx - 12, ty - 7), (tx - 12, ty + 7)], fill="#4B5563")
    for _id, label, x, y in nodes:
        fill = "#E8EEF5" if "Usuario" in label or "App" in label else "#F4F6F9"
        draw.rounded_rectangle((x, y, x + 190, y + 82), radius=12, fill=fill, outline="#2E74B5", width=2)
        lines = wrap_text(draw, label, 165, font(16, True))
        ty = y + 17
        for line in lines[:3]:
            bbox = draw.textbbox((0, 0), line, font=font(16, True))
            draw.text((x + (190 - (bbox[2] - bbox[0])) // 2, ty), line, fill="#111827", font=font(16, True))
            ty += 20
    img.save(out, "PNG")
    return out


def render_mermaid_images() -> list[tuple[str, Path]]:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    outputs = []
    outputs.append((
        "Diagrama Mermaid On-Premise",
        render_flow_diagram(
            "Arquitectura On-Premise",
            [
                ("U", "Usuario web o movil", 60, 190),
                ("K", "Keycloak", 300, 115),
                ("M", "App movil React Native", 300, 275),
                ("G", "API Gateway / Traefik", 550, 195),
                ("S", "Microservicios FastAPI", 790, 195),
                ("A", "Audit Service", 1030, 100),
                ("P", "PostgreSQL", 1030, 215),
                ("N", "MongoDB", 1030, 330),
                ("L", "Loki", 1030, 445),
                ("F", "Grafana", 1260, 445),
            ],
            [("U", "K"), ("U", "M"), ("K", "G"), ("M", "G"), ("G", "S"), ("S", "A"), ("S", "P"), ("S", "N"), ("S", "L"), ("L", "F")],
            OUT_DIR / "mermaid_onprem.png",
        ),
    ))
    outputs.append((
        "Diagrama Mermaid AWS",
        render_flow_diagram(
            "Arquitectura Objetivo AWS",
            [
                ("U", "Usuario web o movil", 80, 235),
                ("CF", "CloudFront / ALB", 330, 235),
                ("ID", "Keycloak / OIDC", 590, 100),
                ("E", "ECS Fargate o EKS", 590, 235),
                ("R", "RDS PostgreSQL", 840, 105),
                ("X", "ElastiCache Redis", 840, 235),
                ("B", "S3", 840, 365),
                ("W", "CloudWatch / Grafana", 1090, 235),
            ],
            [("U", "CF"), ("CF", "ID"), ("CF", "E"), ("E", "R"), ("E", "X"), ("E", "B"), ("E", "W")],
            OUT_DIR / "mermaid_aws.png",
        ),
    ))

    img = Image.new("RGB", (1500, 950), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, 1500, 52), fill="#0B2545")
    draw.text((26, 13), "Mermaid - ERD Logico De Microservicios", fill="#FFFFFF", font=font(22, True))
    boxes = [
        ("CUSTOMERS", 55, 95, "#DAE8FC"),
        ("PRODUCTS", 55, 380, "#D5E8D4"),
        ("ONBOARDING_CASES", 55, 650, "#E1D5E7"),
        ("ACTIVE_SERVICES", 355, 110, "#DAE8FC"),
        ("PROVISIONING_ORDERS", 355, 405, "#F8CECC"),
        ("INVOICES", 690, 110, "#FFE6CC"),
        ("PAYMENTS", 1015, 110, "#FFF2CC"),
        ("NOTIFICATIONS", 690, 430, "#E6D0DE"),
        ("CONSENTS", 360, 700, "#E1D5E7"),
        ("AUDIT_EVENTS", 1015, 430, "#F4CCCC"),
    ]
    dims = {}
    for label, x, y, fill in boxes:
        w, h = 245, 145
        dims[label] = (x, y, w, h)
        draw.rounded_rectangle((x, y, x + w, y + h), radius=12, fill=fill, outline="#4B5563", width=2)
        draw.text((x + 16, y + 14), label, fill="#111827", font=font(17, True))
        sample = {
            "CUSTOMERS": "id PK | document_id | email\nstatus | identity_status",
            "PRODUCTS": "id PK | name | type\nprice | features",
            "ACTIVE_SERVICES": "id PK | customer_id FK\nproduct_id FK | status",
            "PROVISIONING_ORDERS": "id PK | customer_id FK\nproduct_id FK | operation",
            "INVOICES": "id PK | customer_id FK\namount | sri_status",
            "PAYMENTS": "id PK | invoice_id FK\ncustomer_id FK | status",
            "NOTIFICATIONS": "id PK | customer_id FK\nchannel | status",
            "ONBOARDING_CASES": "id PK | document_id\nconsent | kyc_status",
            "CONSENTS": "id PK | user_id FK\nscope | accepted_at",
            "AUDIT_EVENTS": "id PK | service | action\nresource_id | payload JSONB",
        }[label]
        draw.multiline_text((x + 16, y + 54), sample, fill="#374151", font=font(14), spacing=6)

    def connect(a: str, b: str, text: str):
        ax, ay, aw, ah = dims[a]
        bx, by, bw, bh = dims[b]
        sx, sy = ax + aw, ay + ah // 2
        tx, ty = bx, by + bh // 2
        draw.line((sx, sy, tx, ty), fill="#4B5563", width=2)
        draw.polygon([(tx, ty), (tx - 10, ty - 6), (tx - 10, ty + 6)], fill="#4B5563")
        mx, my = (sx + tx) // 2, (sy + ty) // 2 - 18
        draw.text((mx - 35, my), text, fill="#111827", font=font(12, True))

    connect("CUSTOMERS", "ACTIVE_SERVICES", "customer_id")
    connect("PRODUCTS", "ACTIVE_SERVICES", "product_id")
    connect("CUSTOMERS", "INVOICES", "customer_id")
    connect("INVOICES", "PAYMENTS", "invoice_id")
    connect("CUSTOMERS", "NOTIFICATIONS", "customer_id")
    connect("CUSTOMERS", "CONSENTS", "user_id")
    connect("ONBOARDING_CASES", "AUDIT_EVENTS", "resource_id")
    out_erd = OUT_DIR / "mermaid_erd.png"
    img.save(out_erd, "PNG")
    outputs.append(("Diagrama Mermaid ERD", out_erd))
    return outputs


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: RGBColor | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Calibri"
    r.font.size = Pt(9.5)
    if color:
        r.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, h in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, h, bold=True, color=NAVY)
        set_cell_shading(cell, LIGHT_BLUE)
        cell.width = Inches(widths[idx])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
            cells[idx].width = Inches(widths[idx])
    doc.add_paragraph()


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = GRAY


def add_figure(doc: Document, title: str, path: Path, width: float = 6.35) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, title)


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 14, 7),
        ("Heading 2", 13, BLUE, 10, 5),
        ("Heading 3", 11.5, DARK_BLUE, 7, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = "TelcoX | Informe de Arquitectura"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = GRAY
    footer = section.footer.paragraphs[0]
    footer.text = "Documento tecnico generado desde ARCHITECTURE_ONPREM_AWS.md"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8.5)
    footer.runs[0].font.color.rgb = GRAY


def make_report(drawio_images: list[tuple[str, Path]], mermaid_images: list[tuple[str, Path]]) -> None:
    doc = Document()
    configure_doc(doc)
    add_header_footer(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(60)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TelcoX")
    r.font.size = Pt(32)
    r.font.bold = True
    r.font.color.rgb = NAVY
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("Informe Profesional De Arquitectura Aplicada, On-Premise y AWS")
    r.font.size = Pt(18)
    r.font.color.rgb = BLUE
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p3.add_run("Microservicios FastAPI | Next.js | React Native | Kubernetes | CI/CD | Cumplimiento")
    r.font.size = Pt(11)
    r.font.color.rgb = GRAY

    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    items = [
        ("Fuente base", "docs/ARCHITECTURE_ONPREM_AWS.md"),
        ("Diagramas editables", "docs/TELCOX_ARCHITECTURE.drawio"),
        ("Alcance", "Arquitectura implementada, arquitectura AWS objetivo, datos, cumplimiento y DevOps"),
        ("Estado", "Documento de entrega tecnica para revision y presentacion"),
    ]
    for idx, (k, v) in enumerate(items):
        set_cell_text(table.rows[idx].cells[0], k, True, NAVY)
        set_cell_shading(table.rows[idx].cells[0], LIGHT_BLUE)
        set_cell_text(table.rows[idx].cells[1], v)
    doc.add_page_break()

    doc.add_heading("1. Resumen Ejecutivo", level=1)
    doc.add_paragraph(
        "Este informe consolida la arquitectura aplicada en TelcoX y su evolucion objetivo hacia AWS. "
        "La solucion combina microservicios FastAPI, frontend web Next.js, app movil Expo/React Native, "
        "autenticacion OIDC con Keycloak, auditoria centralizada, observabilidad y despliegue Kubernetes con Helm."
    )
    doc.add_paragraph(
        "La arquitectura actual esta preparada para un entorno on-premise pequeno y para una migracion progresiva "
        "a servicios administrados en AWS, manteniendo separacion de responsabilidades por microservicio y trazabilidad operacional."
    )

    add_table(
        doc,
        ["Dimension", "Estado aplicado", "Evidencia"],
        [
            ["Identidad", "Keycloak, OIDC y PKCE", "Login web/movil, tokens de vida corta y bloqueo por onboarding"],
            ["Seguridad", "TLS, cifrado en reposo y Cloudflare WAF", "Controles en Ingress, gateway, frontend y perimetro publico"],
            ["Operacion", "Kubernetes, Helm y observabilidad", "Rolling updates, probes, logs con Loki/Grafana/Promtail"],
            ["DevOps", "CI/CD con GitHub Actions", "Build, push a GHCR, Helm deploy y verificacion de rollout"],
            ["Cumplimiento", "Controles base para telecom y datos sensibles", "Consentimiento, retencion, auditoria y preparacion GDPR/PCI DSS"],
        ],
        [1.45, 2.25, 2.65],
    )

    doc.add_heading("2. Arquitectura Aplicada", level=1)
    doc.add_heading("2.1 Componentes Implementados", level=2)
    for item in [
        "Frontend web con Next.js, login OIDC y barrera de onboarding.",
        "App movil Expo/React Native conectada al mismo flujo OIDC.",
        "Microservicios de customer, catalog, service status, provisioning, billing, payment, notification, onboarding y audit.",
        "Servicios externos simulados para KYC, OSS, pagos, notificaciones y SRI.",
        "Observabilidad con Loki, Grafana y Promtail.",
        "Despliegue Kubernetes con Helm charts por componente.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("2.2 Capturas De Diagramas Draw.io", level=2)
    for idx, (name, path) in enumerate(drawio_images, start=1):
        add_figure(doc, f"Figura {idx}. Captura Draw.io - {name}.", path)

    doc.add_heading("3. Diagramas Mermaid Del Documento", level=1)
    for idx, (name, path) in enumerate(mermaid_images, start=1):
        add_figure(doc, f"Figura M{idx}. {name}.", path)

    doc.add_heading("4. Modelo Logico De Datos", level=1)
    doc.add_paragraph(
        "Los microservicios implementados manejan entidades logicas por dominio. En el entorno de laboratorio, "
        "la mayoria de repositorios son diccionarios en memoria; la infraestructura ya contempla PostgreSQL y MongoDB, "
        "y la auditoria puede persistirse en PostgreSQL mediante AUDIT_BACKEND=postgres."
    )
    add_table(
        doc,
        ["Microservicio", "Entidad logica", "Relaciones principales"],
        [
            ["customer_service", "customers", "Relaciona clientes con servicios activos, facturas, pagos, notificaciones y consentimientos"],
            ["catalog_service", "products", "Relaciona productos con servicios activos y ordenes de provision"],
            ["service_status_service", "active_services", "customer_id y product_id"],
            ["provisioning_service", "provisioning_orders", "customer_id, product_id y referencias de OSS"],
            ["billing_service", "invoices", "customer_id y autorizacion SRI"],
            ["payment_service", "payments", "customer_id, invoice_id y referencia de gateway"],
            ["notification_service", "notifications", "customer_id y gateway externo"],
            ["onboarding_service", "onboarding_cases / consents", "document_id, user_id, KYC y consentimiento explicito"],
            ["audit_service", "audit_events", "service, action, resource_type y resource_id"],
        ],
        [1.65, 1.85, 2.85],
    )

    doc.add_heading("5. Cumplimiento, Seguridad y Trazabilidad", level=1)
    for item in [
        "Ley local de proteccion de datos como marco principal.",
        "Preparacion para GDPR si se tratan datos de usuarios de la Union Europea.",
        "Preparacion para PCI DSS mediante separacion del payment_service y pasarela externa simulada.",
        "TLS en transito, cifrado en reposo y WAF con Cloudflare para el dominio publico.",
        "Consentimiento explicito para documentos, biometria y verificacion de identidad.",
        "Trazabilidad mediante audit_service, logs estructurados y observabilidad centralizada.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("6. Flujo CI/CD DevOps y Extension DevSecOps", level=1)
    doc.add_paragraph(
        "El flujo CI/CD usa GitHub Actions, GitHub Container Registry, Docker, Helm y Kubernetes. "
        "Los pipelines compilan servicios, construyen imagenes, publican tags por commit y despliegan con helm upgrade --install."
    )
    add_table(
        doc,
        ["Pipeline", "Responsabilidad"],
        [
            ["microservices-ci.yml", "Build, push y deploy de microservicios FastAPI"],
            ["ui-ci.yml", "Build de Next.js, imagen Docker y despliegue del frontend"],
            ["audit-service-ci.yml", "Despliegue manual del servicio de auditoria"],
            ["external-services-ci.yml", "Build y despliegue de servicios externos simulados en external-sim"],
        ],
        [2.2, 4.15],
    )
    doc.add_paragraph(
        "La evolucion a DevSecOps puede incorporar SAST, escaneo de dependencias, SBOM, escaneo de imagenes Docker, "
        "validacion de secretos, analisis IaC/Helm/Kubernetes, firmado de imagenes y gates por vulnerabilidades criticas."
    )

    doc.add_heading("7. Recomendaciones De Diseno", level=1)
    for item in [
        "Mantener microservicios desacoplados y evitar compartir logica de dominio entre servicios.",
        "Reutilizar contratos de API antes que modulos compartidos de negocio.",
        "Mantener auditoria separada de observabilidad para preservar finalidad y retencion.",
        "Evolucionar persistencia de demo a PostgreSQL/MongoDB segun criticidad de cada dominio.",
        "Para AWS, priorizar servicios administrados como ECS/EKS, RDS, CloudWatch, Secrets Manager, WAF y ACM.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("8. Conclusion", level=1)
    doc.add_paragraph(
        "TelcoX cuenta con una arquitectura funcional para demo o despliegue inicial on-premise, con una ruta clara hacia AWS. "
        "El diseno actual cubre autenticacion, onboarding, auditoria, observabilidad, despliegue automatizado y controles base de cumplimiento. "
        "La siguiente etapa recomendada es fortalecer persistencia productiva, alta disponibilidad, gobierno de secretos y controles DevSecOps."
    )

    doc.save(OUT_DOCX)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    drawio_images = render_drawio_pages()
    mermaid_images = render_mermaid_images()
    make_report(drawio_images, mermaid_images)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
